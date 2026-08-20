"""Extract plain text from uploaded office/PDF files."""

from __future__ import annotations

import io
import re
import subprocess
import tempfile
from pathlib import Path
from typing import Any

import openpyxl
import pymupdf
import pytesseract
import xlrd
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from loguru import logger
from PIL import Image

from app.core.config import Settings, get_settings

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".doc", ".docx", ".xls", ".xlsx"})

# Numbered list markers: "1.", "2)", "12."
_NUMBERED_ITEM_RE = re.compile(r"^\d+[.)]\s+")
_HEADING_STYLE_RE = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)


class ExtractionError(RuntimeError):
    pass


class TextExtractor:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or get_settings()
        if self.settings.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = self.settings.tesseract_cmd

    def extract(self, filename: str, data: bytes) -> list[str] | list[dict[str, Any]]:
        """Extract text blocks from a file.

        DOCX/DOC return structured blocks::

            {"type": "heading"|"list"|"text"|"table", "content": str, "level": int}

        Other formats return plain ``list[str]``.
        """
        ext = Path(filename).suffix.lower()
        logger.info(
            "Extract start | file={!r} ext={} bytes={}",
            filename,
            ext,
            len(data),
        )
        if ext not in SUPPORTED_EXTENSIONS:
            raise ExtractionError(
                f"Unsupported file type '{ext}'. "
                f"Allowed: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )
        if not data:
            raise ExtractionError("Empty file")

        if ext == ".pdf":
            blocks: list[str] | list[dict[str, Any]] = self._extract_pdf(data)
        elif ext == ".docx":
            blocks = self._extract_docx(data)
        elif ext == ".doc":
            blocks = self._extract_doc(data)
        elif ext == ".xlsx":
            blocks = self._extract_xlsx(data)
        elif ext == ".xls":
            blocks = self._extract_xls(data)
        else:
            raise ExtractionError(f"Unsupported file type '{ext}'")

        logger.info(
            "Extract done | file={!r} blocks={} chars={}",
            filename,
            len(blocks),
            sum(
                len(b if isinstance(b, str) else str(b.get("content", "")))
                for b in blocks
            ),
        )
        return blocks

    def _extract_pdf(self, data: bytes) -> list[str]:
        blocks: list[str] = []
        try:
            doc = pymupdf.open(stream=data, filetype="pdf")
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Failed to open PDF: {exc}") from exc

        with doc:
            logger.debug("PDF pages={}", doc.page_count)
            for page_index, page in enumerate(doc, start=1):
                text = (page.get_text("text") or "").strip()
                if len(text) < self.settings.pdf_ocr_min_chars:
                    logger.info(
                        "PDF page {} low text ({} chars) -> OCR",
                        page_index,
                        len(text),
                    )
                    text = self._ocr_pdf_page(page)
                text = text.strip()
                if text:
                    blocks.append(f"[Страница {page_index}]\n{text}")

        if not blocks:
            raise ExtractionError("No text could be extracted from PDF")
        return blocks

    def _ocr_pdf_page(self, page: pymupdf.Page) -> str:
        try:
            pixmap = page.get_pixmap(dpi=300)
            image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            text = pytesseract.image_to_string(
                image,
                lang=self.settings.ocr_languages,
            )
            logger.debug(
                "OCR page done | langs={} chars={}",
                self.settings.ocr_languages,
                len(text or ""),
            )
            return text
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"OCR failed: {exc}") from exc

    def _extract_docx(self, data: bytes) -> list[dict[str, Any]]:
        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Failed to open DOCX: {exc}") from exc

        parts: list[dict[str, Any]] = []
        table_index = 0

        for block in self._iter_docx_blocks(document):
            if isinstance(block, Paragraph):
                item = self._paragraph_to_block(block)
                if item is not None:
                    parts.append(item)
            elif isinstance(block, Table):
                table_index += 1
                table_text = self._table_to_text(block, table_index)
                if table_text:
                    parts.append(
                        {
                            "type": "table",
                            "content": table_text,
                            "level": 0,
                        }
                    )

        if not parts:
            raise ExtractionError("No text could be extracted from DOCX")
        logger.debug(
            "DOCX parsed | structured blocks={} headings={}",
            len(parts),
            sum(1 for p in parts if p["type"] == "heading"),
        )
        return parts

    @staticmethod
    def _iter_docx_blocks(document: Document):
        """Yield paragraphs and tables in document order."""
        # Prefer modern API when available (python-docx >= 1.1).
        iter_inner = getattr(document, "iter_inner_content", None)
        if callable(iter_inner):
            yield from iter_inner()
            return

        body = document.element.body
        for child in body.iterchildren():
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                yield Paragraph(child, document)
            elif tag == "tbl":
                yield Table(child, document)

    @staticmethod
    def _paragraph_to_block(paragraph: Paragraph) -> dict[str, Any] | None:
        text = (paragraph.text or "").strip()
        if not text:
            return None

        style_name = ""
        try:
            if paragraph.style is not None and paragraph.style.name:
                style_name = str(paragraph.style.name)
        except (AttributeError, ValueError):
            style_name = ""

        heading_match = _HEADING_STYLE_RE.match(style_name)
        if heading_match or style_name.lower() in {"title", "subtitle"}:
            level = int(heading_match.group(1)) if heading_match else 1
            return {"type": "heading", "content": text, "level": level}

        style_lower = style_name.lower()
        is_numbered_style = style_lower.startswith("list number") or (
            "list number" in style_lower
        )
        has_num_pr = TextExtractor._paragraph_has_num_pr(paragraph)
        if (
            _NUMBERED_ITEM_RE.match(text)
            or is_numbered_style
            or has_num_pr
        ):
            return {"type": "list", "content": text, "level": 0}

        return {"type": "text", "content": text, "level": 0}

    @staticmethod
    def _paragraph_has_num_pr(paragraph: Paragraph) -> bool:
        try:
            p_pr = paragraph._element.find(qn("w:pPr"))
            if p_pr is None:
                return False
            return p_pr.find(qn("w:numPr")) is not None
        except Exception:  # noqa: BLE001
            return False

    @staticmethod
    def _table_to_text(table: Table, table_index: int) -> str:
        rows: list[str] = []
        for row in table.rows:
            cells = [
                cell.text.strip() for cell in row.cells if cell.text.strip()
            ]
            if cells:
                rows.append(" | ".join(cells))
        if not rows:
            return ""
        return f"[Таблица {table_index}]\n" + "\n".join(rows)

    def _extract_doc(self, data: bytes) -> list[dict[str, Any]]:
        logger.info(
            "Converting .doc via LibreOffice | path={}",
            self.settings.libreoffice_path,
        )
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "input.doc"
            source.write_bytes(data)
            cmd = [
                self.settings.libreoffice_path,
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                str(tmp_path),
                str(source),
            ]
            try:
                completed = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    timeout=self.settings.libreoffice_timeout_sec,
                    check=False,
                )
            except FileNotFoundError as exc:
                raise ExtractionError(
                    "LibreOffice (soffice) not found. "
                    "Install LibreOffice or set LIBREOFFICE_PATH."
                ) from exc
            except subprocess.TimeoutExpired as exc:
                raise ExtractionError(
                    "LibreOffice conversion timed out"
                ) from exc

            converted = tmp_path / "input.docx"
            if completed.returncode != 0 or not converted.exists():
                stderr = (completed.stderr or completed.stdout or "").strip()
                logger.error(
                    "LibreOffice convert failed | code={} stderr={!r}",
                    completed.returncode,
                    stderr[:500],
                )
                raise ExtractionError(
                    "Failed to convert .doc via LibreOffice"
                    + (f": {stderr}" if stderr else "")
                )
            logger.info(
                "LibreOffice convert ok | bytes={}",
                converted.stat().st_size,
            )
            return self._extract_docx(converted.read_bytes())

    def _extract_xlsx(self, data: bytes) -> list[str]:
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(data),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Failed to open XLSX: {exc}") from exc

        blocks: list[str] = []
        try:
            for sheet in workbook.worksheets:
                lines: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    cells = [
                        str(cell).strip()
                        for cell in row
                        if cell is not None and str(cell).strip()
                    ]
                    if cells:
                        lines.append(" | ".join(cells))
                if lines:
                    blocks.append(f"[Лист {sheet.title}]\n" + "\n".join(lines))
                    logger.debug(
                        "XLSX sheet={!r} rows={}",
                        sheet.title,
                        len(lines),
                    )
        finally:
            workbook.close()

        if not blocks:
            raise ExtractionError("No text could be extracted from XLSX")
        return blocks

    def _extract_xls(self, data: bytes) -> list[str]:
        try:
            book = xlrd.open_workbook(file_contents=data)
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Failed to open XLS: {exc}") from exc

        blocks: list[str] = []
        for sheet in book.sheets():
            lines: list[str] = []
            for row_index in range(sheet.nrows):
                cells = [
                    str(sheet.cell_value(row_index, col_index)).strip()
                    for col_index in range(sheet.ncols)
                    if str(sheet.cell_value(row_index, col_index)).strip()
                ]
                if cells:
                    lines.append(" | ".join(cells))
            if lines:
                blocks.append(f"[Лист {sheet.name}]\n" + "\n".join(lines))
                logger.debug("XLS sheet={!r} rows={}", sheet.name, len(lines))

        if not blocks:
            raise ExtractionError("No text could be extracted from XLS")
        return blocks
